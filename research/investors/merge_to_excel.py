#!/usr/bin/env python3
"""Merge worker JSONL results into the Investors sheet as new columns.

KEY: Company ID is NOT unique in this data (one placeholder ID is shared by 291
distinct investors), so we match by INVESTOR NAME restricted to the target set
(Current Holdings > 0 AND Transaction in Last 24 Months = Yes). Names are unique
within that set (except one harmless 2x case), so this never smears a profile
onto an unrelated row.

Adds four columns on the Investors sheet:
  Ankaufsprofil (recherchiert) | Investor-Website | Quelle | Recherche-Status
"""
import json, glob, os, sys
import openpyxl

SRC = '/root/.claude/uploads/ca96adda-b743-5c4f-9e92-3404803593a5/59067e43-Properties_and_Investors_1.xlsx'
OUT = sys.argv[1] if len(sys.argv) > 1 else 'research/investors/Properties_and_Investors_ANKAUFSPROFIL.xlsx'
WORKDIR = 'research/investors/workers'

def norm(s):
    return ' '.join(str(s).strip().split()).lower() if s is not None else ''

# 1) collect worker results by normalized investor name (last write wins)
results = {}
n_lines = 0
for fp in sorted(glob.glob(os.path.join(WORKDIR, 'worker-*.jsonl'))):
    for line in open(fp, encoding='utf-8'):
        line = line.strip()
        if not line:
            continue
        try:
            rec = json.loads(line)
        except Exception as e:
            print(f'  ! skip bad json in {fp}: {e}', file=sys.stderr)
            continue
        key = norm(rec.get('investor', ''))
        if not key:
            continue
        results[key] = rec
        n_lines += 1
print(f'Loaded {n_lines} records, {len(results)} unique investor-name keys')

# 2) load workbook (writable)
wb = openpyxl.load_workbook(SRC)
ws = wb['Investors']
hdr = [c.value for c in ws[1]]
name_col = 1  # 'Investor' is column A
h_col = hdr.index('Current Holdings (objects)') + 1
l_col = hdr.index('Transaction in Last 24 Months') + 1

new_cols = ['Ankaufsprofil (recherchiert)', 'Investor-Website', 'Quelle', 'Recherche-Status']
base = ws.max_column
for i, name in enumerate(new_cols):
    ws.cell(row=1, column=base + 1 + i, value=name)
c_prof, c_web, c_src, c_stat = (base + 1, base + 2, base + 3, base + 4)

filled = 0
matched_keys = set()
for r in range(2, ws.max_row + 1):
    nm = ws.cell(row=r, column=name_col).value
    if nm is None:
        continue
    h = ws.cell(row=r, column=h_col).value or 0
    l = ws.cell(row=r, column=l_col).value
    if not (h > 0 and l == 'Yes'):   # only target rows
        continue
    rec = results.get(norm(nm))
    if not rec:
        continue
    ws.cell(row=r, column=c_prof, value=rec.get('ankaufsprofil', ''))
    ws.cell(row=r, column=c_web, value=rec.get('website', ''))
    ws.cell(row=r, column=c_src, value=rec.get('quelle', ''))
    ws.cell(row=r, column=c_stat, value=rec.get('status', ''))
    filled += 1
    matched_keys.add(norm(nm))

unmatched = sorted(set(results) - matched_keys)
os.makedirs(os.path.dirname(OUT), exist_ok=True)
wb.save(OUT)
print(f'Filled {filled} target rows. Wrote {OUT}')
if unmatched:
    print(f'WARNING: {len(unmatched)} worker records did not match a target row by name:')
    for k in unmatched[:30]:
        print('   -', results[k].get('investor'))
